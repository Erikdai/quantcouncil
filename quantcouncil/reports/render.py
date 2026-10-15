"""Render a Report to Word (python-docx) or PDF (ReportLab, built-in CJK font)."""

from __future__ import annotations

import io
import re
from functools import lru_cache
from xml.sax.saxutils import escape

from quantcouncil.reports.document import Bullets, Heading, Para, Report, Table

ACCENT = (0x36, 0x4F, 0xC7)
MUTED = (0x6B, 0x72, 0x80)


# --- Word ------------------------------------------------------------------------------

def to_docx(report: Report) -> bytes:
    from docx import Document  # noqa: PLC0415
    from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: PLC0415
    from docx.oxml.ns import qn  # noqa: PLC0415
    from docx.shared import Pt, RGBColor  # noqa: PLC0415

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for level in range(1, 5):
        style = doc.styles[f"Heading {level}"]
        style.font.color.rgb = RGBColor(*ACCENT)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    def runs(paragraph, items, muted=False):
        for value, bold in items:
            run = paragraph.add_run(value)
            run.bold = bold
            if muted:
                run.font.color.rgb = RGBColor(*MUTED)
                run.font.size = Pt(9)

    doc.add_heading(report.title, level=0)
    runs(doc.add_paragraph(), [(report.subtitle, False)], muted=True)
    for block in report.blocks:
        if isinstance(block, Heading):
            doc.add_heading(block.text, level=block.level)
        elif isinstance(block, Para):
            runs(doc.add_paragraph(), block.runs, muted=block.muted)
        elif isinstance(block, Bullets):
            for item in block.items:
                runs(doc.add_paragraph(style="List Bullet"), item)
        elif isinstance(block, Table):
            has_header = any(block.header)
            cols = max(len(block.header), *(len(r) for r in block.rows)) if block.rows else len(block.header)
            table = doc.add_table(rows=0, cols=cols)
            table.style, table.alignment = "Light Grid Accent 1", WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate(([block.header] if has_header else []) + block.rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=False):
                    cell.text = value
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9)
                            r.bold = has_header and i == 0
            doc.add_paragraph()
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- PDF -------------------------------------------------------------------------------

@lru_cache(maxsize=1)
def _styles():
    from reportlab.lib.colors import Color  # noqa: PLC0415
    from reportlab.lib.fonts import addMapping  # noqa: PLC0415
    from reportlab.lib.styles import ParagraphStyle  # noqa: PLC0415
    from reportlab.pdfbase import pdfmetrics  # noqa: PLC0415
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # noqa: PLC0415

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))  # ships with ReportLab; no font files needed
    for bold in (0, 1):
        addMapping("STSong-Light", bold, 0, "STSong-Light")  # no bold CJK face: <b> keeps the glyphs
    accent = Color(*(c / 255 for c in ACCENT))
    muted = Color(*(c / 255 for c in MUTED))
    base = ParagraphStyle("base", fontName="STSong-Light", fontSize=10, leading=15, wordWrap="CJK")
    return {
        "title": ParagraphStyle("title", parent=base, fontSize=20, leading=26, textColor=accent, spaceAfter=4),
        "subtitle": ParagraphStyle("subtitle", parent=base, fontSize=9, textColor=muted, spaceAfter=12),
        1: ParagraphStyle("h1", parent=base, fontSize=15, leading=20, textColor=accent, spaceBefore=12, spaceAfter=6),
        2: ParagraphStyle("h2", parent=base, fontSize=12.5, leading=17, textColor=accent, spaceBefore=9, spaceAfter=4),
        3: ParagraphStyle("h3", parent=base, fontSize=11, leading=15, spaceBefore=6, spaceAfter=3),
        4: ParagraphStyle("h4", parent=base, fontSize=10.5, leading=14, spaceBefore=4, spaceAfter=2),
        "body": ParagraphStyle("body", parent=base, spaceAfter=5),
        "muted": ParagraphStyle("muted", parent=base, fontSize=8.5, leading=12, textColor=muted, spaceAfter=5),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8.5, leading=11),
        "accent": accent,
    }


_CJK = re.compile(r"([⺀-鿿　-〿＀-￯]+)")


def _markup(items) -> str:
    """ReportLab markup: CJK in the CID font, everything else in Helvetica (the CID font's Latin is poor)."""
    out = []
    for value, bold in items:
        for i, part in enumerate(_CJK.split(value)):
            if not part:
                continue
            if i % 2:  # a CJK run
                out.append(f"<b>{escape(part)}</b>" if bold else escape(part))
            else:
                face = "Helvetica-Bold" if bold else "Helvetica"
                out.append(f'<font name="{face}">{escape(part)}</font>')
    return "".join(out).replace("\n", "<br/>")


def to_pdf(report: Report) -> bytes:
    from reportlab.lib import colors  # noqa: PLC0415
    from reportlab.lib.pagesizes import A4  # noqa: PLC0415
    from reportlab.lib.units import mm  # noqa: PLC0415
    from reportlab.platypus import (  # noqa: PLC0415
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table as PdfTable,
        TableStyle,
    )

    s = _styles()
    story = [Paragraph(_markup([(report.title, False)]), s["title"]),
             Paragraph(_markup([(report.subtitle, False)]), s["subtitle"])]
    width = A4[0] - 36 * mm
    for block in report.blocks:
        if isinstance(block, Heading):
            story.append(Paragraph(_markup([(block.text, False)]), s[block.level]))
        elif isinstance(block, Para):
            story.append(Paragraph(_markup(block.runs), s["muted" if block.muted else "body"]))
        elif isinstance(block, Bullets):
            story.append(ListFlowable([ListItem(Paragraph(_markup(i), s["body"]), leftIndent=12)
                                       for i in block.items], bulletType="bullet", bulletFontSize=6, leftIndent=12))
        elif isinstance(block, Table):
            rows = ([block.header] if any(block.header) else []) + block.rows
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            data = [[Paragraph(_markup([(c, False)]), s["cell"]) for c in r + [""] * (cols - len(r))] for r in rows]
            table = PdfTable(data, colWidths=[width / cols] * cols, repeatRows=1 if any(block.header) else 0)
            style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d0d7e2")),
                     ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fb")])]
            if any(block.header):
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7ecfb")))
            table.setStyle(TableStyle(style))
            story += [table, Spacer(1, 8)]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawString(18 * mm, 10 * mm, report.title)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(doc.page))
        canvas.restoreState()

    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                      bottomMargin=18 * mm, title=report.title).build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


MEDIA_TYPES = {"pdf": "application/pdf",
               "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


def render(report: Report, fmt: str) -> bytes:
    return to_pdf(report) if fmt == "pdf" else to_docx(report)
